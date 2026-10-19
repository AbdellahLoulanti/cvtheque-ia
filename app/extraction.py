"""Extraction et nettoyage du texte des CV (PDF, DOCX, TXT)."""

from __future__ import annotations

import io
import re
import unicodedata
from pathlib import Path

import pymupdf
from docx import Document

from app.config import ACCEPTED_CV_EXTENSIONS, MIN_TEXT_LETTERS


class ExtractionError(Exception):
    """Le fichier ne peut pas être lu (format non supporté, corrompu, protégé…)."""


_TXT_ENCODINGS = ("utf-8-sig", "cp1252", "latin-1")

_INVISIBLE_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\u00ad\u200b-\u200f\u2028\u2029\ufeff]")
# Puces Word/PDF encodées dans la zone Unicode privée (ex. \uf0b7)
_PRIVATE_USE_CHARS = re.compile(r"[\ue000-\uf8ff]")
_HORIZONTAL_SPACES = re.compile(r"[ \t\f\v\u00a0\u2000-\u200a\u202f\u205f\u3000]+")
_EXTRA_BLANK_LINES = re.compile(r"\n{3,}")


def extract_text(filename: str, data: bytes) -> str:
    """Retourne le texte brut (non nettoyé) d'un CV selon son extension."""
    ext = Path(filename).suffix.lower()
    if ext not in ACCEPTED_CV_EXTENSIONS:
        raise ExtractionError(f"Format non supporté : {ext or 'sans extension'}")

    try:
        if ext == ".pdf":
            return _extract_pdf(data)
        if ext == ".docx":
            return _extract_docx(data)
        return _extract_txt(data)
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001 — fichier corrompu, format invalide…
        raise ExtractionError(f"Lecture impossible : {exc}") from exc


def _extract_pdf(data: bytes) -> str:
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        if doc.needs_pass:
            raise ExtractionError("PDF protégé par mot de passe")
        return "\n".join(page.get_text() for page in doc)


def _extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    # Beaucoup de CV sont mis en page avec des tableaux
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(parts)


def _extract_txt(data: bytes) -> str:
    for encoding in _TXT_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Encodage du fichier texte non reconnu")


def clean_text(text: str) -> str:
    """Normalise les espaces, retours à la ligne et supprime les caractères parasites."""
    text = unicodedata.normalize("NFC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _INVISIBLE_CHARS.sub("", text)
    text = _PRIVATE_USE_CHARS.sub("•", text)
    text = _HORIZONTAL_SPACES.sub(" ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    text = _EXTRA_BLANK_LINES.sub("\n\n", text)
    return text.strip()


def has_usable_text(text: str) -> bool:
    """Faux pour un fichier vide ou un PDF scanné (image sans couche texte)."""
    return sum(char.isalpha() for char in text) >= MIN_TEXT_LETTERS
